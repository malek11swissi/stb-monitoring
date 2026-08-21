from pathlib import Path
from textwrap import wrap
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "report"
ASSETS = OUT / "assets"
OUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

NAVY = "123B5D"
TEAL = "087E8B"
CYAN = "2CB5C0"
LIGHT = "EAF1F4"
GRAY = "D9D9D9"
DARK = "18252E"
MUTED = "61727E"
RED = "D64545"
GREEN = "198754"
ORANGE = "E89A19"
PURPLE = "7257A8"

def font(size, bold=False):
    paths = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for p in paths:
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()

def rounded_box(draw, xy, title, lines, fill, accent=None, width=3):
    x1,y1,x2,y2=xy
    outline = accent or NAVY
    if not outline.startswith("#"):
        outline = "#" + outline
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=width)
    draw.text((x1+18,y1+14), title, font=font(25,True), fill="#"+NAVY)
    y=y1+52
    for line in lines:
        draw.text((x1+20,y), line, font=font(18), fill="#"+DARK)
        y+=28

def arrow(draw, start, end, label=None):
    draw.line([start,end], fill="#"+TEAL, width=5)
    import math
    a=math.atan2(end[1]-start[1],end[0]-start[0])
    p1=(end[0]-18*math.cos(a-.55),end[1]-18*math.sin(a-.55))
    p2=(end[0]-18*math.cos(a+.55),end[1]-18*math.sin(a+.55))
    draw.polygon([end,p1,p2], fill="#"+TEAL)
    if label:
        mx=(start[0]+end[0])//2; my=(start[1]+end[1])//2
        draw.text((mx-55,my-25),label,font=font(16,True),fill="#"+MUTED)

def save_canvas(name, title, size=(1800,1050)):
    im=Image.new("RGB",size,"white")
    d=ImageDraw.Draw(im)
    d.text((70,42),title,font=font(36,True),fill="#"+NAVY)
    d.line((70,95,size[0]-70,95),fill="#"+CYAN,width=5)
    return im,d,ASSETS/name

def make_figures():
    im,d,p=save_canvas("architecture-generale.png","Architecture générale de STB Sentinel")
    rounded_box(d,(80,210,390,510),"Utilisateurs",["Administrateur","Superviseur","Technicien","Manager IT"],"#F4F8FA")
    rounded_box(d,(555,160,930,560),"STB Sentinel",["Frontend Angular","API centrale .NET 8","Workers planifiés","PostgreSQL"],"#E7F4F5",TEAL)
    rounded_box(d,(1110,140,1710,580),"SI surveillés",["Core Banking / Oracle","RNE / MongoDB","SMS / MySQL","RH / SQL Server"],"#F4F8FA")
    rounded_box(d,(555,700,930,930),"Services externes",["Brevo - e-mail","Twilio - SMS"],"#F6F2FA",PURPLE)
    arrow(d,(390,360),(555,360),"HTTPS")
    arrow(d,(930,360),(1110,360),"HTTP/TLS")
    arrow(d,(745,560),(745,700),"API")
    im.save(p)

    im,d,p=save_canvas("chaine-supervision.png","Chaîne active de supervision et de réaction")
    boxes=[("Planifier",["NextCheckAt","Verrou distribué"]),("Contrôler",["HTTP","API JSON","TLS"]),("Historiser",["CheckResult","Preuve technique"]),("Évaluer",["Règle","Déduplication"]),("Agir",["Alerte","Incident","Notification"])]
    xs=[60,410,760,1110,1460]
    for i,(t,ls) in enumerate(boxes):
        rounded_box(d,(xs[i],260,xs[i]+280,590),t,ls,"#F4F8FA",TEAL)
        if i<len(boxes)-1: arrow(d,(xs[i]+280,425),(xs[i+1],425))
    d.text((220,760),"UP",font=font(28,True),fill="#"+GREEN)
    d.text((500,760),"DEGRADED",font=font(28,True),fill="#"+ORANGE)
    d.text((930,760),"DOWN",font=font(28,True),fill="#"+RED)
    d.text((1310,760),"UNKNOWN",font=font(28,True),fill="#"+MUTED)
    im.save(p)

    im,d,p=save_canvas("architecture-logique.png","Architecture logique en couches")
    layers=[("Présentation",["Angular","Pages par rôle","Guards et formulaires"],"#E7F4F5"),("API / Application",["Contrôleurs REST","Cas d'utilisation","Autorisation"],"#EEF3F8"),("Domaine",["SI et endpoints","Alertes et incidents","SLA et transitions"],"#F4F8FA"),("Infrastructure",["EF Core / PostgreSQL","HTTP, TLS","Brevo et Twilio"],"#F6F2FA")]
    y=150
    for t,ls,c in layers:
        rounded_box(d,(250,y,1550,y+180),t,ls,c,TEAL)
        y+=215
    im.save(p)

    im,d,p=save_canvas("microservices-simules.png","Écosystème de simulation des SI")
    data=[("Core Banking","Oracle","5101 / 7101"),("RNE","MongoDB","5102 / 7102"),("SMS","MySQL","5103 / 7103"),("RH","SQL Server","5104 / 7104")]
    x=70
    for t,db,port in data:
        rounded_box(d,(x,230,x+390,650),t,["API .NET 8",db,"HTTP / HTTPS",port],"#F4F8FA",TEAL)
        x+=435
    d.text((250,800),"Chaque simulateur est démarrable, testable et conteneurisable indépendamment.",font=font(27,True),fill="#"+NAVY)
    im.save(p)

    im,d,p=save_canvas("validation-tls.png","Validation réelle d'un certificat TLS")
    rounded_box(d,(70,220,380,610),"Endpoint",["URL HTTPS","Nom DNS","Port 443"],"#F4F8FA")
    rounded_box(d,(520,180,900,650),"TLS handshake",["Connexion TCP","Certificat présenté","TLS 1.2 / 1.3"],"#E7F4F5",TEAL)
    rounded_box(d,(1040,130,1720,700),"Contrôles",["Dates de validité","Nom d'hôte","Chaîne de confiance","Expiration proche","Connexion / négociation"],"#F4F8FA")
    arrow(d,(380,410),(520,410))
    arrow(d,(900,410),(1040,410))
    d.text((280,820),"La preuve enregistrée provient du certificat réellement présenté par le serveur.",font=font(27,True),fill="#"+NAVY)
    im.save(p)

    im,d,p=save_canvas("cycle-incident.png","Cycle de vie opérationnel d'un incident")
    states=[("New",NAVY),("Assigned",TEAL),("InProgress",CYAN),("Pending",ORANGE),("Resolved",GREEN),("Closed",MUTED)]
    x=60
    for i,(s,c) in enumerate(states):
        d.rounded_rectangle((x,310,x+245,500),radius=18,fill="#F4F8FA",outline="#"+c,width=5)
        tw=d.textbbox((0,0),s,font=font(24,True))[2]
        d.text((x+(245-tw)//2,385),s,font=font(24,True),fill="#"+c)
        if i<len(states)-1: arrow(d,(x+245,405),(x+285,405))
        x+=285
    d.text((190,700),"Superviseur : qualification et affectation",font=font(24,True),fill="#"+NAVY)
    d.text((670,760),"Technicien : diagnostic et résolution",font=font(24,True),fill="#"+TEAL)
    d.text((1160,700),"Validation et clôture",font=font(24,True),fill="#"+GREEN)
    im.save(p)

    im,d,p=save_canvas("classes-globales.png","Agrégats métier et relations principales")
    rounded_box(d,(80,180,390,460),"User",["Rôle fixe","Préférences","Audit"],"#F4F8FA")
    rounded_box(d,(500,150,830,490),"MonitoredSystem",["Environnement","Criticité","État global"],"#E7F4F5",TEAL)
    rounded_box(d,(940,150,1270,490),"Endpoint",["HTTP / API / TLS","Fréquence","Seuils"],"#F4F8FA")
    rounded_box(d,(1380,150,1710,490),"CheckResult",["Statut","Durée","Preuve"],"#F4F8FA")
    rounded_box(d,(500,650,830,930),"Alert",["Règle","Déduplication","Occurrences"],"#FFF7E8",ORANGE)
    rounded_box(d,(940,630,1270,950),"Incident",["Affectation","SLA","Résolution"],"#FCEEEE",RED)
    rounded_box(d,(1380,650,1710,930),"Notification",["Interne","E-mail","SMS"],"#F6F2FA",PURPLE)
    arrow(d,(390,320),(500,320)); arrow(d,(830,320),(940,320)); arrow(d,(1270,320),(1380,320))
    arrow(d,(1545,490),(665,650)); arrow(d,(830,790),(940,790)); arrow(d,(1270,790),(1380,790))
    im.save(p)

    im,d,p=save_canvas("devops-cible.png","Chaîne DevOps cible")
    steps=[("GitHub",["Code source"]),("Jenkins",["Build","Tests"]),("Registry",["Images Docker"]),("Kubernetes",["Déploiement","Probes"]),("Observabilité",["Prometheus","Grafana"])]
    xs=[55,400,745,1090,1435]
    for i,(t,ls) in enumerate(steps):
        rounded_box(d,(xs[i],270,xs[i]+280,600),t,ls,"#F4F8FA",TEAL)
        if i<len(steps)-1: arrow(d,(xs[i]+280,435),(xs[i+1],435))
    d.text((310,770),"Cible d'industrialisation - les manifestes Kubernetes et pipelines restent à finaliser.",font=font(25,True),fill="#"+NAVY)
    im.save(p)

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"),fill)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=tcMar.find(qn("w:"+m))
        if node is None: node=OxmlElement("w:"+m); tcMar.append(node)
        node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa")

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement("w:tblHeader"); el.set(qn("w:val"),"true"); trPr.append(el)

def add_field(paragraph, code):
    r=paragraph.add_run(); begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin")
    instr=OxmlElement("w:instrText"); instr.set(qn("xml:space"),"preserve"); instr.text=code
    sep=OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"),"separate")
    txt=OxmlElement("w:t"); txt.text="Mettre à jour le champ dans Word"
    end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
    r._r.extend([begin,instr,sep,txt,end])

def style_document(doc):
    sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.0); sec.left_margin=Cm(2.5); sec.right_margin=Cm(2.3)
    sec.header_distance=Cm(.8); sec.footer_distance=Cm(.8)
    normal=doc.styles["Normal"]
    normal.font.name="Times New Roman"; normal.font.size=Pt(11.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"),"Times New Roman"); normal._element.rPr.rFonts.set(qn("w:hAnsi"),"Times New Roman")
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.2
    for name,size,color,before,after in (("Heading 1",18,NAVY,18,10),("Heading 2",15,NAVY,14,7),("Heading 3",12.5,MUTED,10,5)):
        s=doc.styles[name]; s.font.name="Arial"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
        s._element.rPr.rFonts.set(qn("w:ascii"),"Arial"); s._element.rPr.rFonts.set(qn("w:hAnsi"),"Arial")
        s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
    cap=doc.styles["Caption"]; cap.font.name="Times New Roman"; cap.font.size=Pt(10); cap.font.italic=True; cap.font.color.rgb=RGBColor.from_string(DARK)
    cap.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after=Pt(8)

def running_header_footer(section):
    hp=section.header.paragraphs[0]; hp.clear(); hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=hp.add_run("STB Sentinel - Rapport de projet de fin d'études"); r.font.name="Times New Roman"; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(MUTED)
    pPr=hp._p.get_or_add_pPr(); pbdr=OxmlElement("w:pBdr"); bot=OxmlElement("w:bottom"); bot.set(qn("w:val"),"single"); bot.set(qn("w:sz"),"6"); bot.set(qn("w:color"),MUTED); pbdr.append(bot); pPr.append(pbdr)
    fp=section.footer.paragraphs[0]; fp.clear(); fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=fp.add_run("Page "); rr.font.size=Pt(9); add_field(fp,"PAGE")

def add_cover(doc):
    sec=doc.sections[0]; sec.different_first_page_header_footer=True
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(30)
    r=p.add_run("ESPRIT"); r.font.name="Arial"; r.font.size=Pt(30); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(RED)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("2025 - 2026\nPROJET DE FIN D'ÉTUDES"); r.font.name="Arial"; r.font.size=Pt(18); r.font.bold=True
    band=doc.add_table(rows=1,cols=1); band.alignment=WD_TABLE_ALIGNMENT.CENTER; band.autofit=False; band.columns[0].width=Cm(16)
    set_cell_shading(band.cell(0,0),NAVY); bp=band.cell(0,0).paragraphs[0]; bp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    br=bp.add_run("DIPLÔME NATIONAL D'INGÉNIEUR"); br.font.name="Arial"; br.font.size=Pt(18); br.font.bold=True; br.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("STB SENTINEL"); r.font.name="Arial"; r.font.size=Pt(28); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Plateforme proactive de supervision et de gestion des incidents des systèmes d'information bancaires"); r.font.name="Arial"; r.font.size=Pt(19); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(DARK)
    doc.add_paragraph()
    meta=[("Réalisé par", "Malek Jendoubi"),("Encadrant académique", "Monsieur Hamdi Braiek"),("Organisme d'étude", "Société Tunisienne de Banque - contexte académique simulé"),("Spécialité", "TWIN")]
    t=doc.add_table(rows=len(meta),cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,(a,b) in enumerate(meta):
        t.cell(i,0).width=Cm(5); t.cell(i,1).width=Cm(10)
        t.cell(i,0).text=a; t.cell(i,1).text=b
        for c in t.rows[i].cells:
            set_cell_margins(c,110,160,110,160); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in c.paragraphs[0].runs: run.font.name="Arial"; run.font.size=Pt(11)
        t.cell(i,0).paragraphs[0].runs[0].bold=True; set_cell_shading(t.cell(i,0),LIGHT)
    doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Version de rédaction fondée sur l'état réel du dépôt au 20 août 2026"); r.font.name="Arial"; r.font.size=Pt(9.5); r.font.italic=True; r.font.color.rgb=RGBColor.from_string(MUTED)
    doc.add_page_break()

def add_front_page(doc,title,paras):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(40); p.paragraph_format.space_after=Pt(35)
    set_cell=None
    r=p.add_run(title); r.font.name="Arial"; r.font.size=Pt(20); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
    for text in paras:
        q=doc.add_paragraph(text); q.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_page_break()

def add_chapter(doc,n,title,intro):
    doc.add_page_break()
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(65); p.paragraph_format.space_after=Pt(24)
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),GRAY); pPr.append(shd)
    r=p.add_run(f"{n}  {title}"); r.font.name="Arial"; r.font.size=Pt(22); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(DARK)
    doc.add_heading("Introduction",level=1); doc.add_paragraph(intro)

def add_table(doc, headers, rows, widths=None, caption=None):
    if caption:
        p=doc.add_paragraph(caption,style="Caption")
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    set_repeat_table_header(t.rows[0])
    total=16.2
    if widths is None: widths=[total/len(headers)]*len(headers)
    for i,h in enumerate(headers):
        c=t.cell(0,i); c.text=str(h); c.width=Cm(widths[i]); set_cell_shading(c,NAVY); set_cell_margins(c)
        for r in c.paragraphs[0].runs: r.font.name="Arial"; r.font.size=Pt(9); r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].width=Cm(widths[i]); set_cell_margins(cells[i])
            if len(t.rows)%2==0: set_cell_shading(cells[i],"F5F8FA")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
                for r in p.runs: r.font.name="Arial"; r.font.size=Pt(8.5)
    return t

def add_figure(doc, filename, caption, width=15.8):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS/filename),width=Cm(width))
    doc.add_paragraph(caption,style="Caption")

def bullet(doc,text):
    p=doc.add_paragraph(style="List Bullet"); p.add_run(text)

def numbered(doc,text):
    p=doc.add_paragraph(style="List Number"); p.add_run(text)

def build():
    make_figures()
    doc=Document(); style_document(doc); running_header_footer(doc.sections[0]); add_cover(doc)
    add_front_page(doc,"Dédicaces",["À mes parents, pour leurs sacrifices, leur patience et leur soutien constant tout au long de mon parcours.","À ma famille, à mes amis et à toutes les personnes qui m'ont encouragé et accompagné dans la réalisation de ce projet.","Je leur dédie humblement ce travail."])
    add_front_page(doc,"Remerciements",["Au terme de ce projet de fin d'études, je tiens à exprimer ma gratitude à toutes les personnes qui ont contribué à son aboutissement.","J'adresse mes sincères remerciements à mon encadrant académique, Monsieur Hamdi Braiek, pour ses conseils, sa disponibilité et son accompagnement. Je remercie également le corps professoral et administratif d'ESPRIT pour la qualité de la formation dispensée.","Enfin, je remercie les membres du jury pour l'attention portée à ce travail ainsi que toutes les personnes qui m'ont soutenu durant sa réalisation."])
    add_front_page(doc,"Résumé",["La continuité des systèmes d'information constitue un enjeu majeur pour les établissements bancaires. Une indisponibilité, une dégradation progressive ou l'expiration d'un certificat numérique peut affecter les opérations internes et la qualité du service rendu aux clients. Ce projet propose STB Sentinel, une plateforme proactive de supervision et de gestion des incidents conçue dans un contexte académique représentatif de la Société Tunisienne de Banque.","La solution centralise la configuration des systèmes et de leurs points de contrôle, exécute des vérifications HTTP, API JSON et TLS, calcule les états UP, DEGRADED, DOWN et UNKNOWN, puis transforme les anomalies confirmées en alertes dédupliquées et en incidents actionnables. Elle couvre également l'affectation aux techniciens, les SLA, les escalades, les maintenances, les notifications internes, e-mail et SMS, ainsi que les indicateurs MTTD, MTTA et MTTR et les exports PDF, Excel et CSV.","La plateforme repose sur Angular, ASP.NET Core .NET 8, Entity Framework Core et PostgreSQL. Quatre microservices indépendants simulent le Core Banking, le RNE, le service SMS et le système RH avec respectivement Oracle, MongoDB, MySQL et SQL Server. Cette approche permet de valider les mécanismes de supervision avant tout raccordement aux SI réels.","Mots-clés : supervision, systèmes bancaires, disponibilité, TLS, alertes, incidents, SLA, Angular, .NET 8, PostgreSQL, microservices, DevOps."])
    doc.add_heading("Table des matières",level=1); p=doc.add_paragraph(); add_field(p,'TOC \\o "1-3" \\h \\z \\u'); doc.add_page_break()
    doc.add_heading("Liste des abréviations",level=1)
    abbr=[("API","Application Programming Interface"),("CI/CD","Continuous Integration / Continuous Delivery"),("CRUD","Create, Read, Update, Delete"),("DNS","Domain Name System"),("HTTP","HyperText Transfer Protocol"),("HTTPS","HTTP Secure"),("JWT","JSON Web Token"),("KPI","Key Performance Indicator"),("MTTA","Mean Time To Acknowledge"),("MTTD","Mean Time To Detect"),("MTTR","Mean Time To Resolve"),("RBAC","Role-Based Access Control"),("REST","Representational State Transfer"),("SI","Système d'information"),("SLA","Service Level Agreement"),("TLS","Transport Layer Security"),("UML","Unified Modeling Language")]
    add_table(doc,["Abréviation","Signification"],abbr,[4.2,12.0]); doc.add_page_break()

    doc.add_heading("Introduction générale",level=1)
    for ptxt in [
        "La transformation numérique du secteur bancaire augmente la dépendance aux applications, aux API et aux infrastructures interconnectées. La défaillance d'un seul composant peut perturber une chaîne de traitement, retarder une opération ou dégrader l'expérience des collaborateurs et des clients.",
        "Une supervision efficace ne consiste pas seulement à vérifier qu'un serveur répond. Elle doit mesurer la disponibilité, la latence, la conformité des réponses métier et la sécurité des échanges HTTPS. Elle doit également transformer les observations techniques en alertes pertinentes, puis en incidents suivis par les équipes selon des responsabilités et des délais clairement établis.",
        "STB Sentinel répond à cette problématique au moyen d'une plateforme centrale de monitoring actif. En l'absence d'accès aux SI réels durant la phase académique, quatre simulateurs reproduisent des comportements représentatifs et permettent d'injecter des anomalies contrôlées. Le présent rapport expose le contexte, l'analyse des besoins, la méthodologie Scrum, la conception, la réalisation et la trajectoire DevOps de la solution."
    ]: doc.add_paragraph(ptxt)

    add_chapter(doc,1,"Cadre général du projet","Ce chapitre présente le contexte bancaire étudié, la problématique de continuité des SI, les limites d'une surveillance fragmentée et la solution proposée.")
    doc.add_heading("1.1 Présentation du contexte d'accueil",level=1)
    doc.add_paragraph("La Société Tunisienne de Banque évolue dans un environnement où les applications métier, les services de communication, les référentiels externes et les systèmes de ressources humaines participent quotidiennement aux opérations. Dans ce rapport, aucune donnée ni infrastructure réelle de la banque n'est exposée : les noms fonctionnels servent à construire un laboratoire académique contrôlé.")
    doc.add_heading("1.2 Contexte et problématique",level=1)
    doc.add_paragraph("La multiplication des dépendances applicatives rend difficile une détection rapide des pannes. Un service peut répondre tout en étant incapable d'accéder à sa base de données ; un endpoint peut rester disponible mais devenir trop lent ; un certificat TLS peut expirer et interrompre brutalement les échanges HTTPS. Sans centralisation, la détection repose souvent sur des vérifications manuelles ou sur les réclamations des utilisateurs.")
    doc.add_heading("1.3 Objectifs",level=1)
    for x in ["centraliser le catalogue des SI et de leurs environnements","exécuter des contrôles HTTP, API JSON et TLS configurables","conserver une preuve technique de chaque contrôle","dédupliquer les anomalies et corréler les alertes","piloter le cycle complet des incidents et des SLA","notifier les acteurs appropriés","produire des tableaux de bord et rapports exploitables","préparer un déploiement reproductible et observable"]: bullet(doc,x)
    doc.add_heading("1.4 Étude de l'existant",level=1)
    add_table(doc,["Approche","Atouts","Limites dans le contexte"],[
        ("Vérification manuelle","Simple au démarrage","Non continue, non historisée, dépendante d'une personne"),
        ("Logs applicatifs isolés","Détail technique riche","Vision fragmentée et réaction tardive"),
        ("Outil générique de disponibilité","Mesures et tableaux de bord","Peu de gestion métier des incidents sans intégration"),
        ("STB Sentinel","Contrôle actif, preuve, alerte, incident et SLA","Nécessite des endpoints, un accès réseau et une configuration fiable")], [4,5.8,6.4], "Tableau 1.1 - Comparaison des approches de supervision")
    doc.add_heading("1.5 Solution proposée",level=1)
    doc.add_paragraph("STB Sentinel adopte une plateforme centrale modulaire. Elle interroge les endpoints depuis l'extérieur des SI, mesure leur comportement et applique des règles de santé homogènes. Les simulateurs sont indépendants de la plateforme et reproduisent les technologies de données retenues pour les quatre écosystèmes.")
    add_figure(doc,"architecture-generale.png","Figure 1.1 - Architecture générale de STB Sentinel")
    doc.add_heading("Conclusion",level=1); doc.add_paragraph("La solution vise à réduire le délai de détection et à structurer la réaction opérationnelle, tout en restant adaptable à de futurs SI réels.")

    add_chapter(doc,2,"Analyse des besoins et méthodologie","Ce chapitre identifie les acteurs, formalise les besoins fonctionnels et non fonctionnels, puis organise le projet selon Scrum.")
    doc.add_heading("2.1 Identification des acteurs",level=1)
    actors=[("Administrateur","Comptes, catalogue SI, endpoints, règles et configuration technique"),("Superviseur","Alertes, qualification, création et affectation des incidents"),("Technicien","Diagnostic et résolution des incidents qui lui sont affectés"),("Manager IT","KPI, SLA, rapports, risques et validation des décisions critiques"),("SI surveillé","Expose les URLs HTTP/HTTPS et certificats contrôlés"),("Fournisseur externe","Envoie les e-mails ou SMS demandés par la plateforme")]
    add_table(doc,["Acteur","Responsabilité"],actors,[4.2,12.0],"Tableau 2.1 - Acteurs de STB Sentinel")
    doc.add_heading("2.2 Besoins fonctionnels",level=1)
    needs=[("Identité","Connexion JWT, profil, comptes actifs, rôles fixes et audit"),("Catalogue","SI, environnement, criticité, propriétaire et archivage"),("Endpoints","URL, type, fréquence, timeout, seuils, code et JSON attendus"),("Supervision","Planification, exécution manuelle, état et historique"),("TLS","Chaîne, dates, hôte, confiance, négociation et preuve"),("Alertes","Règles, confirmation, déduplication, acquittement et résolution"),("Incidents","Corrélation, affectation, timeline, résolution et pièces jointes"),("SLA","Échéances, AtRisk, Breached, Met et escalade"),("Maintenance","Calendrier jour/semaine/mois et suspension d'alertes"),("Reporting","KPI, filtres, vues, PDF, Excel et CSV")]
    add_table(doc,["Gestion","Besoins couverts"],needs,[4.0,12.2],"Tableau 2.2 - Synthèse des besoins fonctionnels")
    doc.add_heading("2.3 Besoins non fonctionnels",level=1)
    for x in ["Sécurité : authentification, autorisation côté API, mots de passe hachés et secrets non exposés.","Fiabilité : annulation, timeouts, verrou distribué et conservation des preuves.","Performance : contrôles planifiés indépendants, pagination et index PostgreSQL.","Disponibilité : endpoints liveness/readiness et services conteneurisables.","Traçabilité : audit, historique d'incident, horodatages et résultats d'envoi.","Maintenabilité : séparation Domain, Application, Infrastructure et API.","Ergonomie : interface responsive, états lisibles et messages orientés action."]: bullet(doc,x)
    doc.add_heading("2.4 Répartition des privilèges",level=1)
    perms=[("Utilisateurs","Gérer","Lire","-","Lire"),("SI / endpoints","Gérer","Lire et tester","Lire le périmètre","Lire"),("Alertes","Lire","Acquitter et qualifier","Liées aux incidents","Lire les critiques"),("Incidents","Lire","Créer/affecter/piloter","Traiter les siens","Valider les critiques"),("SLA / rapports","Configurer techniquement","Suivre","Consulter le sien","Piloter et valider"),("Maintenance","Administrer","Planifier","Consulter","Approuver production")]
    add_table(doc,["Gestion","Admin","Superviseur","Technicien","Manager IT"],perms,[3.1,3.1,3.5,3.3,3.2],"Tableau 2.3 - Matrice cible des responsabilités")
    doc.add_heading("2.5 Méthodologie Scrum",level=1)
    doc.add_paragraph("Le projet est découpé en incréments courts. Chaque sprint possède un objectif, des User Stories, des critères d'acceptation et une démonstration. Le Product Backlog est priorisé par valeur métier et risque technique.")
    backlog=[("Sprint 0","Cadrage et environnement"),("Sprint 1","Identité et accès"),("Sprint 2","Catalogue SI et endpoints"),("Sprint 3","Monitoring HTTP/API/TLS"),("Sprint 4","Simulateurs et bases indépendantes"),("Sprint 5","Alertes et incidents"),("Sprint 6","Notifications, SLA et escalades"),("Sprint 7","Maintenance, dashboard et rapports"),("Sprint 8","Sécurité et qualité"),("Sprint 9","DevOps et observabilité"),("Sprint 10","Stabilisation et soutenance")]
    add_table(doc,["Sprint","Objectif"],backlog,[3.0,13.2],"Tableau 2.4 - Planification synthétique des sprints")
    doc.add_heading("2.6 User Stories structurantes",level=1)
    stories=[("US-101","Utilisateur","Se connecter et gérer son profil"),("US-201","Administrateur","Créer un SI et ses endpoints"),("US-301","Superviseur","Lancer un contrôle manuel"),("US-303","Plateforme","Détecter disponibilité, latence et certificat"),("US-501","Plateforme","Dédupliquer les alertes"),("US-502","Superviseur","Qualifier et affecter un incident"),("US-503","Technicien","Documenter et résoudre son incident"),("US-601","Utilisateur","Recevoir des notifications internes, e-mail et SMS"),("US-701","Manager IT","Consulter santé, SLA et rapports")]
    add_table(doc,["ID","Acteur","Besoin"],stories,[2.4,3.7,10.1],"Tableau 2.5 - User Stories principales")
    doc.add_heading("2.7 Definition of Done",level=1)
    for x in ["code compilé et relu","tests adaptés au risque","autorisation API et visibilité frontend cohérentes","migration fournie si le modèle change","documentation mise à jour","démonstration reproductible","aucun secret commité"]: bullet(doc,x)
    doc.add_heading("Conclusion",level=1); doc.add_paragraph("Le découpage Scrum rend la progression démontrable et permet de traiter en priorité la chaîne de valeur allant de la détection à la résolution.")

    add_chapter(doc,3,"Conception","Ce chapitre transforme les besoins en architecture, modèles métier, relations de données et scénarios d'interaction.")
    doc.add_heading("3.1 Architecture générale",level=1); doc.add_paragraph("La solution sépare l'interface Angular, l'API centrale, les workers et les adaptateurs d'infrastructure. Les SI simulés restent externes au cœur de supervision et sont interrogés via HTTP/HTTPS.")
    add_figure(doc,"architecture-logique.png","Figure 3.1 - Architecture logique en couches")
    doc.add_heading("3.2 Choix architectural",level=1)
    doc.add_paragraph("Le backend STB Sentinel est un monolithe modulaire en couches, et non un ensemble de microservices internes. Ce choix limite la complexité transactionnelle tout en maintenant une séparation nette des responsabilités. Les quatre SI simulés sont, eux, des services déployables indépendamment avec leurs propres bases.")
    add_figure(doc,"microservices-simules.png","Figure 3.2 - Microservices simulant les SI")
    doc.add_heading("3.3 Communication",level=1)
    doc.add_paragraph("HTTP/HTTPS constitue le mécanisme principal car le monitoring doit observer le système depuis l'extérieur. RabbitMQ n'est pas obligatoire pour qualifier une architecture distribuée et ne remplace pas le polling lorsqu'un SI est totalement indisponible. Il est conservé comme évolution possible pour des événements asynchrones tels que AlertCreated ou IncidentResolved.")
    doc.add_heading("3.4 Modèle métier",level=1)
    add_figure(doc,"classes-globales.png","Figure 3.3 - Agrégats métier et relations principales")
    doc.add_paragraph("Le modèle persistant comprend actuellement dix-huit tables métier. Les énumérations sont stockées sous forme de texte ; le rôle est un champ de User et aucune table roles, permissions ou user_roles n'est nécessaire.")
    db=[("users","Comptes et rôles"),("monitored_systems","SI et état agrégé"),("monitoring_endpoints","Configuration des contrôles"),("check_results","Historique et preuves"),("alert_rules","Règles dynamiques"),("alerts / alert_occurrences","Anomalies dédupliquées"),("incidents","Cycle de traitement"),("incident_*","Commentaires, histoire, preuves et escalade"),("sla_policies","Objectifs de service"),("notifications","Flux interne"),("maintenance_windows","Périodes planifiées"),("saved_views","Filtres utilisateurs")]
    add_table(doc,["Table ou groupe","Rôle"],db,[5.2,11.0],"Tableau 3.1 - Organisation conceptuelle des données")
    doc.add_heading("3.5 Conception du monitoring",level=1)
    add_figure(doc,"chaine-supervision.png","Figure 3.4 - Chaîne de supervision")
    doc.add_paragraph("L'ordre d'enregistrement est important : le résultat technique est persisté avant l'évaluation des alertes, ce qui garantit qu'une alerte référence une preuve déjà disponible.")
    doc.add_heading("3.6 Conception TLS",level=1)
    add_figure(doc,"validation-tls.png","Figure 3.5 - Processus de validation TLS")
    doc.add_paragraph("Le contrôle TLS ne dépend pas de données de certificat saisies par l'administrateur. À chaque exécution, le moteur ouvre une connexion réelle, récupère le certificat présenté, construit la chaîne et enregistre les métadonnées. En production, la CA interne de la STB devra être installée dans le magasin de confiance du serveur ou du conteneur.")
    doc.add_heading("3.7 Cycle d'incident",level=1)
    add_figure(doc,"cycle-incident.png","Figure 3.6 - Machine d'état d'un incident")
    doc.add_heading("Conclusion",level=1); doc.add_paragraph("La conception isole les règles métier des dépendances techniques et prépare le remplacement des simulateurs par des URLs réelles sans recompilation.")

    add_chapter(doc,4,"Réalisation","Ce chapitre décrit les fonctionnalités effectivement présentes dans le dépôt et explique leur contribution à la chaîne métier.")
    doc.add_heading("4.1 Identité, authentification et audit",level=1)
    doc.add_paragraph("L'API authentifie l'utilisateur à partir de son nom ou de son e-mail, vérifie le mot de passe haché et l'état du compte, puis génère un JWT contenant le rôle. Les tentatives réussies ou échouées sont auditées. La logique empêche également la désactivation ou le changement de rôle du dernier administrateur actif.")
    doc.add_heading("4.2 Catalogue des SI et environnements",level=1)
    doc.add_paragraph("Un SI est défini par son code, son nom, sa description, son environnement, sa criticité et son propriétaire. Il peut être désactivé ou archivé. L'archivage est privilégié à la suppression lorsqu'un historique existe, afin de préserver les preuves et rapports.")
    envs=[("Production","Système exploité et critique"),("Preproduction","Validation avant mise en production"),("Recette","Tests fonctionnels"),("Development","Développement et laboratoire")]
    add_table(doc,["Environnement","Usage"],envs,[4.3,11.9],"Tableau 4.1 - Environnements gérés")
    doc.add_heading("4.3 Configuration des endpoints",level=1)
    endpoint_fields=[("URL","Adresse absolue HTTP/HTTPS"),("CheckType","Http, ApiJson ou Tls"),("HttpMethod","GET, HEAD ou POST"),("ExpectedStatusCode","Code attendu, généralement 200"),("TimeoutSeconds","Durée maximale d'attente"),("IntervalSeconds","Période entre deux contrôles"),("DegradedThresholdMs","Seuil de latence dégradée"),("DownThresholdMs","Seuil critique"),("IsCritical","Impact sur l'état global du SI"),("ExpectedJsonProperty / Value","Validation métier de la réponse")]
    add_table(doc,["Champ","Interprétation"],endpoint_fields,[5.1,11.1],"Tableau 4.2 - Champs principaux d'un endpoint")
    doc.add_heading("4.4 Exécution des contrôles",level=1)
    doc.add_paragraph("MonitoringWorker interroge les endpoints arrivés à échéance. Un verrou consultatif PostgreSQL empêche plusieurs instances de traiter simultanément le même cycle. Chaque endpoint choisit un exécuteur spécialisé et l'échec d'un contrôle n'empêche pas la poursuite des autres contrôles.")
    doc.add_heading("4.4.1 Contrôle HTTP",level=2); doc.add_paragraph("Le moteur vérifie l'accessibilité, le code de réponse, le timeout et la latence. Une réponse correcte peut donc être classée DEGRADED ou DOWN si elle dépasse les seuils configurés.")
    doc.add_heading("4.4.2 Contrôle API JSON",level=2); doc.add_paragraph("Le contrôle API étend HTTP par la lecture d'une propriété JSON. Il permet par exemple de vérifier database.available=true. Cette stratégie détecte une base indisponible uniquement si l'endpoint contrôlé dépend réellement de cette base ou expose explicitement son état.")
    doc.add_heading("4.4.3 Contrôle TLS",level=2); doc.add_paragraph("TlsCheckExecutor utilise TcpClient et SslStream pour obtenir le certificat réel. Il distingue expiration, date future, nom incorrect, chaîne non approuvée, expiration proche, connexion impossible et négociation impossible. La chaîne complète est sérialisée dans les métadonnées. La révocation en ligne CRL/OCSP reste à activer pour une cible bancaire de production.")
    tls_cases=[("TLS_EXPIRED","Certificat expiré","DOWN"),("TLS_EXPIRING","Moins de 30 jours","DEGRADED ou DOWN"),("TLS_NAME_MISMATCH","Nom DNS absent du certificat","DOWN"),("TLS_UNTRUSTED","Chaîne non approuvée","DOWN"),("TLS_HANDSHAKE","Négociation impossible","DOWN"),("TLS_CONNECTION","Connexion impossible","DOWN")]
    add_table(doc,["Code","Signification","État"],tls_cases,[4.0,8.5,3.7],"Tableau 4.3 - Classification des erreurs TLS")
    doc.add_heading("4.5 Alertes",level=1)
    doc.add_paragraph("Les règles d'alerte sont des entités persistées et administrables. Le moteur attend le nombre d'échecs configuré, construit une clé de déduplication et réutilise l'alerte active pendant la fenêtre définie. Chaque répétition reste liée à un CheckResult.")
    stdrules=[("Endpoint indisponible","2","15 min"),("Performance dégradée","2","15 min"),("Timeout réseau","1","15 min"),("Réponse API invalide","1","15 min"),("Certificat bientôt expiré","1","24 h"),("Certificat expiré","1","24 h"),("Certificat invalide","1","60 min")]
    add_table(doc,["Règle standard","Échecs","Déduplication"],stdrules,[8.2,3.4,4.6],"Tableau 4.4 - Règles initiales")
    doc.add_heading("4.6 Incidents",level=1)
    doc.add_paragraph("Le superviseur peut créer un incident depuis une alerte. Une recherche sur le même SI et une fenêtre récente permet de rattacher plusieurs alertes à une panne commune. Le technicien documente le résumé, la cause racine, l'action corrective, l'action préventive et la preuve de résolution.")
    doc.add_heading("4.7 SLA et escalade",level=1)
    sla=[("P1 Critical","15 min","60 min"),("P2 High","30 min","240 min"),("P3 Medium","120 min","480 min"),("P4 Low","480 min","1440 min")]
    add_table(doc,["Priorité","Réponse","Résolution"],sla,[5.2,5.5,5.5],"Tableau 4.5 - Politiques SLA initiales")
    doc.add_paragraph("Un worker recalcule les états OnTrack, AtRisk, Breached et Met. Pour les P1, l'escalade persistée notifie successivement le superviseur puis le manager IT si l'incident reste ouvert. La persistance permet de reprendre le traitement après un redémarrage.")
    doc.add_heading("4.8 Notifications",level=1)
    doc.add_paragraph("Les notifications internes sont stockées par utilisateur. ApiNotificationChannel utilise Brevo pour l'e-mail et Twilio pour le SMS, avec un mode Simulation qui permet de tester le contrat sans consommer de quota ni stocker de secrets réels.")
    doc.add_heading("4.9 Maintenance",level=1)
    doc.add_paragraph("Le calendrier propose les vues jour, semaine et mois, ainsi que création, modification, annulation et suppression. Une fenêtre peut conserver les contrôles tout en supprimant temporairement la création d'alertes. Les maintenances annulées sont exclues du calendrier actif.")
    doc.add_heading("4.10 Dashboard et rapports",level=1)
    doc.add_paragraph("Le dashboard agrège santé des SI, disponibilité, alertes, incidents, SLA et tendances. Les mesures MTTD, MTTA et MTTR représentent respectivement les délais moyens de détection, d'acquittement et de résolution. Les résultats sont exportables en CSV, PDF et Excel .xlsx sur des périodes mensuelles ou annuelles.")
    metrics=[("MTTD","Temps entre début estimé et détection"),("MTTA","Temps entre alerte et acquittement"),("MTTR","Temps entre création et résolution"),("Disponibilité","Part des contrôles réussis"),("Conformité SLA","Part des incidents résolus dans le délai")]
    add_table(doc,["Indicateur","Interprétation"],metrics,[4.0,12.2],"Tableau 4.6 - Indicateurs opérationnels")
    doc.add_heading("4.11 Interface Angular",level=1)
    ui=[("Connexion","Authentification et retour vers la route demandée"),("Dashboard","Synthèse adaptée au rôle"),("Systèmes","Catalogue et filtres"),("Détail SI","Endpoints, contrôles et preuves TLS"),("Alertes","Liste, acquittement et règles"),("Incidents","Filtres, vues sauvegardées et affectation"),("Détail incident","Timeline, pièces jointes et résolution"),("Calendrier","Maintenances jour/semaine/mois"),("Profil","Identité et préférences de notification")]
    add_table(doc,["Écran","Finalité"],ui,[4.3,11.9],"Tableau 4.7 - Principaux écrans Angular")
    doc.add_heading("4.12 Limites et travaux en cours",level=1)
    limits=[("Privilèges","La matrice cible de séparation stricte doit remplacer les droits encore trop larges de l'Admin."),("TLS production","Activer révocation en ligne et éventuellement mTLS."),("RabbitMQ","Mentionné comme cible mais absent de l'implémentation actuelle."),("DevOps","Docker des simulateurs présent ; Jenkins/Kubernetes/Prometheus/Grafana restent à finaliser."),("Accès SI réels","Nécessite DNS, firewall, endpoints autorisés et CA interne de confiance.")]
    add_table(doc,["Sujet","État ou action"],limits,[4.2,12.0],"Tableau 4.8 - Écarts entre implémentation et cible")
    doc.add_heading("Conclusion",level=1); doc.add_paragraph("La réalisation couvre le cœur fonctionnel de la chaîne de supervision et rend chaque résultat traçable. Les travaux restants concernent principalement le durcissement des rôles, la production TLS et l'industrialisation.")

    add_chapter(doc,5,"DevOps, tests et déploiement","Ce chapitre présente les éléments disponibles dans le dépôt et la cible d'industrialisation retenue.")
    doc.add_heading("5.1 Tests automatisés",level=1)
    tests=[("Unitaires","User, sprints 1-3, maintenance et cycle incident"),("Architecture","Respect des dépendances entre couches"),("Autorisation","Réponses 401/403 et politiques"),("Intégration PostgreSQL","Migrations, persistance et cycle métier")]
    add_table(doc,["Famille","Couverture"],tests,[4.3,11.9],"Tableau 5.1 - Organisation des tests")
    doc.add_paragraph("La solution compile actuellement sans erreur ni avertissement. Les tests doivent évoluer avec la nouvelle matrice des privilèges afin de vérifier que chaque rôle ne peut exécuter que ses responsabilités.")
    doc.add_heading("5.2 Conteneurisation",level=1)
    doc.add_paragraph("Chaque simulateur possède un Dockerfile et un fichier Docker Compose associé à sa base. La base PostgreSQL centrale est également démarrable par Compose. Cette organisation permet d'arrêter une dépendance isolément et d'observer l'impact depuis STB Sentinel.")
    doc.add_heading("5.3 Cible CI/CD et Kubernetes",level=1)
    add_figure(doc,"devops-cible.png","Figure 5.1 - Chaîne DevOps cible")
    doc.add_paragraph("La cible prévoit une image par composant, des manifestes Deployment, Service, ConfigMap, Secret et PVC, des probes liveness/readiness et des rolling updates. Jenkins doit enchaîner compilation, tests, analyse, construction des images, publication et smoke tests. Ces éléments sont décrits comme perspectives tant que leurs fichiers ne sont pas présents dans le dépôt.")
    doc.add_heading("5.4 Observabilité",level=1)
    doc.add_paragraph("Prometheus et Grafana doivent compléter la supervision fonctionnelle par des métriques techniques portant sur l'API, les workers, les simulateurs et les pods. Les métriques recommandées comprennent le nombre et la durée des contrôles, les erreurs, les alertes créées, les notifications échouées et l'état des exécutions planifiées.")
    doc.add_heading("5.5 Sécurité de déploiement",level=1)
    for x in ["injection des secrets par variables protégées ou Kubernetes Secrets","bases non exposées publiquement","CA interne montée dans les conteneurs de confiance","images versionnées par SHA Git","probes et limites de ressources","sauvegarde et procédure de rollback","aucune donnée réelle dans le laboratoire académique"]: bullet(doc,x)
    doc.add_heading("Conclusion",level=1); doc.add_paragraph("La structure actuelle prépare l'industrialisation, mais la chaîne DevOps complète doit être finalisée et démontrée avant de la présenter comme acquise.")

    doc.add_page_break(); doc.add_heading("Conclusion générale et perspectives",level=1)
    for ptxt in [
        "STB Sentinel démontre qu'une plateforme centrale peut couvrir une chaîne cohérente allant du contrôle technique à la gestion opérationnelle. L'utilisation d'endpoints configurables permet de surveiller disponibilité, performance, contenu métier et certificats sans lier le moteur à un SI particulier.",
        "Le recours à quatre simulateurs indépendants a permis d'éprouver le comportement de la plateforme en l'absence d'accès aux systèmes réels. Cette séparation reste fidèle à une intégration future : le passage en environnement STB reposera principalement sur le remplacement des URLs, l'ouverture des flux réseau, la gestion des authentifications et l'installation des autorités de certification internes.",
        "Les perspectives prioritaires sont le durcissement de la séparation des privilèges, la validation de révocation TLS, la prise en charge éventuelle de mTLS, l'industrialisation Kubernetes/Jenkins, l'observabilité Prometheus/Grafana et l'ajout ultérieur d'une détection prédictive explicable."
    ]: doc.add_paragraph(ptxt)
    doc.add_heading("Bilan",level=2); doc.add_paragraph("Le projet a permis de mettre en œuvre une architecture en couches, un moteur de monitoring actif, des règles dynamiques, une gestion d'incidents structurée et une stratégie de test couvrant domaine, autorisation et persistance.")
    doc.add_heading("Difficultés rencontrées",level=2)
    for x in ["modéliser des SI non accessibles réellement","distinguer disponibilité d'API et santé de base de données","valider TLS sans accepter silencieusement un certificat incorrect","éviter les doubles exécutions et tempêtes d'alertes","définir une séparation cohérente entre rôles techniques et décisionnels"]: bullet(doc,x)
    doc.add_heading("Perspectives",level=2)
    for x in ["connexion contrôlée aux SI réels","révocation TLS et mTLS","service d'événements asynchrones si un besoin RabbitMQ est confirmé","déploiement Kubernetes reproductible","alerting technique Prometheus/Grafana","modèles d'anomalie explicables avec validation humaine"]: bullet(doc,x)

    doc.add_page_break(); doc.add_heading("Bibliographie et webographie",level=1)
    refs=[
        "Schwaber, K. et Sutherland, J. The Scrum Guide, version 2020. https://scrumguides.org/scrum-guide.html",
        "Microsoft. Documentation ASP.NET Core et .NET 8. https://learn.microsoft.com/aspnet/core/",
        "Angular. Documentation officielle. https://angular.dev/docs",
        "PostgreSQL Global Development Group. Documentation PostgreSQL. https://www.postgresql.org/docs/",
        "Kubernetes. Concepts et documentation officielle. https://kubernetes.io/docs/concepts/",
        "OWASP Foundation. Application Security Verification Standard. https://owasp.org/www-project-application-security-verification-standard/",
        "IETF. RFC 8446 - The Transport Layer Security (TLS) Protocol Version 1.3. https://www.rfc-editor.org/rfc/rfc8446"
    ]
    for i,x in enumerate(refs,1): doc.add_paragraph(f"{i}. {x}")
    doc.add_heading("Annexe A - Commandes principales",level=1)
    cmds=[("Backend","dotnet run --project platform/backend/StbMonitoring.Api/StbMonitoring.Api.csproj"),("Frontend","npm start depuis platform/frontend"),("Tests","dotnet test StbMonitoring.sln"),("RNE","docker compose -f simulators/docker-compose.rne.yml up --build"),("SMS","docker compose -f simulators/docker-compose.sms.yml up --build"),("RH","docker compose -f simulators/docker-compose.hr.yml up --build"),("Core Banking","docker compose -f simulators/docker-compose.core-banking.yml up --build")]
    add_table(doc,["Composant","Commande"],cmds,[4.0,12.2],"Tableau A.1 - Commandes de démarrage et test")
    doc.add_heading("Annexe B - Scénario de démonstration",level=1)
    demo=["Démarrer PostgreSQL, l'API, Angular et un simulateur.","Ajouter ou vérifier le SI et ses endpoints HTTP/API/TLS.","Exécuter un contrôle et constater l'état UP.","Arrêter la base du simulateur ou activer une anomalie.","Attendre les échecs consécutifs et observer l'alerte dédupliquée.","Acquitter l'alerte, créer l'incident et l'affecter.","Documenter la résolution et relancer le contrôle.","Vérifier le retour UP, la résolution et le rapport."]
    for i,x in enumerate(demo,1): doc.add_paragraph(f"{i}. {x}")

    for s in doc.sections: running_header_footer(s)
    props=doc.core_properties; props.title="STB Sentinel - Rapport PFE"; props.author="Malek Jendoubi"; props.subject="Supervision et gestion des incidents des SI bancaires"; props.keywords="STB Sentinel, monitoring, TLS, incidents, SLA"
    path=OUT/"Rapport_PFE_STB_Sentinel_Malek_Jendoubi.docx"; doc.save(path); print(path)

if __name__=="__main__": build()
